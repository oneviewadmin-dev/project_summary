"""PPT 파일을 글자로 바꿔 읽는다.

PowerPoint 파일(.pptx)을 열어 슬라이드마다 제목 · 본문 · 표 · 이슈 · 참고를
뽑아내고, 마지막에 AI 가 읽기 좋은 글 한 덩어리로 합친다.

  .pptx 파일 → PPTXParser.parse() → ParsedPPT → to_llm_input() → AI 에게 전달
"""
from pptx import Presentation
from dataclasses import dataclass, field
from typing import List


# ==========================================
#  슬라이드 안에 있던 표 한 개
# ==========================================
@dataclass
class SlideTable:
    headers: List[str]        # 표의 첫 줄 (머리글)
    rows: List[List[str]]     # 나머지 줄들

    # ==========================================
    #  표를 글자로 바꾼다. AI 는 그림을 못 보고 글자만 읽기 때문이다.
    #  (마크다운 = 글자로 표를 그리는 방식. | 로 칸을 나눈다)
    # ==========================================
    def to_markdown(self) -> str:
        if not self.headers:
            return ""
        lines = ["| " + " | ".join(self.headers) + " |"]
        lines.append("|" + "|".join(["---"] * len(self.headers)) + "|")
        for r in self.rows:
            lines.append("| " + " | ".join(r) + " |")
        return "\n".join(lines)


# ==========================================
#  슬라이드 한 장에서 꺼낸 것
# ==========================================
@dataclass
class SlideContent:
    slide_number: int                                      # 몇 번째 장인가
    title: str                                             # 그 장의 제목
    body_text: str                                         # 본문 글
    tables: List[SlideTable] = field(default_factory=list)  # 그 장에 있던 표들
    issues: List[str] = field(default_factory=list)         # "※ 이슈사항" 아래 줄들
    references: List[str] = field(default_factory=list)     # "※ 참고사항" 아래 줄들


# ==========================================
#  PPT 파일 한 개를 통째로 읽은 결과
# ==========================================
@dataclass
class ParsedPPT:
    file_name: str
    author: str                  # 표지에 적힌 작성자
    date: str                    # 표지에 적힌 작성일
    title: str                   # 표지에 적힌 제목
    slides: List[SlideContent]   # 장별 내용

    # ==========================================
    #  읽은 내용을 AI 에게 넘길 글 한 덩어리로 합친다.
    #  장마다 제목 → 본문 → 표 → 이슈 → 참고 순서로 이어 붙인다.
    # ==========================================
    def to_llm_input(self) -> str:
        out = [f"# {self.title}", f"작성자: {self.author} / 작성일: {self.date}\n"]
        for s in self.slides:
            out.append(f"\n## [슬라이드 {s.slide_number}] {s.title}")
            if s.body_text:
                out.append(s.body_text)
            for t in s.tables:
                out.append("\n" + t.to_markdown())
            if s.issues:
                out.append("\n[이슈사항]\n- " + "\n- ".join(s.issues))
            if s.references:
                out.append("\n[참고사항]\n- " + "\n- ".join(s.references))
        return "\n".join(out)


# ==========================================
#  실제로 PPT 를 여는 쪽. 위 세 가지를 채워서 돌려준다.
# ==========================================
class PPTXParser:
    # 이 말이 나오면 그 아래 줄부터는 '이슈' 로 모은다
    ISSUE_MARKERS = ["※ 이슈사항", "이슈사항", "리스크"]
    # 이 말이 나오면 그 아래 줄부터는 '참고' 로 모은다
    REF_MARKERS = ["※ 참고사항", "참고사항", "특이사항"]

    # ==========================================
    #  바깥에서 부르는 입구. 파일 경로를 받아 읽은 결과를 돌려준다.
    # ==========================================
    def parse(self, path: str) -> ParsedPPT:
        prs = Presentation(path)
        slides = [self._parse_slide(s, i + 1) for i, s in enumerate(prs.slides)]
        # 표지(첫 장)에서 제목 · 작성자 · 작성일을 따로 꺼낸다
        meta = self._parse_cover(prs.slides[0]) if prs.slides else {}
        return ParsedPPT(
            file_name=path.split("/")[-1],
            author=meta.get("author", ""),
            date=meta.get("date", ""),
            title=meta.get("title", ""),
            slides=slides,
        )

    # ==========================================
    #  표지 한 장에서 제목 · 작성자 · 작성일을 찾는다.
    #  정해진 자리가 없어서, 글자를 보고 짐작한다.
    # ==========================================
    def _parse_cover(self, slide) -> dict:
        # ----------------------------------------------------------------------------
        #  표지에 있는 글상자의 글을 전부 이어 붙인다
        # ----------------------------------------------------------------------------
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        full = "\n".join(texts)

        # ----------------------------------------------------------------------------
        #  줄마다 보면서 "작성일" · "작성자" 가 적힌 줄을 가려낸다.
        #  둘 다 아니면서 처음 나오는 줄을 제목으로 본다
        #  ('confidential' 같은 보안 문구는 제목이 아니므로 건너뛴다)
        # ----------------------------------------------------------------------------
        meta = {"title": "", "author": "", "date": ""}
        for line in full.split("\n"):
            line = line.strip()
            if not line:
                continue
            if "작성일" in line:
                meta["date"] = line.replace("작성일", "").strip()
            elif "작성자" in line:
                meta["author"] = line.replace("작성자", "").strip()
            elif not meta["title"] and "confidential" not in line.lower():
                meta["title"] = line
        return meta

    # ==========================================
    #  슬라이드 한 장을 읽어 제목 · 본문 · 표 · 이슈 · 참고로 나눈다
    # ==========================================
    def _parse_slide(self, slide, num: int) -> SlideContent:
        title = ""
        body, issues, refs, tables = [], [], [], []

        for shape in slide.shapes:
            # ----------------------------------------------------------------------------
            #  표면 표대로 읽는다
            # ----------------------------------------------------------------------------
            if shape.has_table:
                tables.append(self._parse_table(shape.table))

            elif shape.has_text_frame:
                text = shape.text_frame.text

                # ----------------------------------------------------------------------------
                #  제목 고르기 — 짧고(60자 미만) 한 줄짜리 글을 제목으로 본다.
                #  한 장에 하나만 잡는다
                # ----------------------------------------------------------------------------
                if not title and 0 < len(text.strip()) < 60 and "\n" not in text:
                    title = text.strip()
                    continue

                # ----------------------------------------------------------------------------
                #  본문 나누기 — "※ 이슈사항" 같은 말이 나오면 거기서부터
                #  이슈로, "※ 참고사항" 이 나오면 참고로 모은다.
                #  아무것도 안 나왔으면 그냥 본문이다
                # ----------------------------------------------------------------------------
                section = None
                for line in text.split("\n"):
                    s = line.strip()
                    if not s:
                        continue
                    if any(m in s for m in self.ISSUE_MARKERS):
                        section = "issue"
                        continue
                    if any(m in s for m in self.REF_MARKERS):
                        section = "ref"
                        continue
                    if section == "issue":
                        issues.append(s)
                    elif section == "ref":
                        refs.append(s)
                    else:
                        body.append(s)

        return SlideContent(
            slide_number=num,
            title=title,
            body_text="\n".join(body),
            tables=tables,
            issues=issues,
            references=refs,
        )

    # ==========================================
    #  표 한 개를 읽는다. 첫 줄은 머리글, 나머지는 내용
    # ==========================================
    def _parse_table(self, tbl) -> SlideTable:
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
        return SlideTable(headers=rows[0] if rows else [], rows=rows[1:])
