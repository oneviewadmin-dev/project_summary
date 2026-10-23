# 어느 프로젝트든 양식의 열 자리가 같다
#   python3 backend/tests/test_form_same_columns.py
#
# 사업부 여덟 곳이 제각각인 양식을 쓰고 있어서 전용 파서가 11종까지 늘었다.
# 하나로 맞추려면 "엑셀을 열었을 때 늘 같은 자리에 같은 것이 있다" 가
# 먼저다. 전에는 그 프로젝트가 안 쓰는 칸을 빼서, 챔버는 파트넘버가
# A열인데 다른 프로젝트는 모델명이 A열이었다.
#
# 이 시험이 막는 것
#   프로젝트마다 열을 빼거나 더해서 자리가 다시 어긋나는 것.
import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                      # noqa: E402
import plan_form                                        # noqa: E402
import weekly_template as wt                            # noqa: E402

ok = 0

# ── 1. 모델이 어떻게 생겼든 왼쪽 칸은 넷이다 ─────────────────
WANT = ["파트넘버", "모델명", "유형", "구분"]
CASES = {
    "파트넘버·유형 다 있음": [{"id": "a", "name": "A", "part_number": "P1",
                               "dev_type": "413", "group": "양산"}],
    "파트넘버 없음":         [{"id": "b", "name": "B", "group": "양산"}],
    "유형 없음":             [{"id": "c", "name": "C", "part_number": "P3",
                               "group": "개발"}],
    "이름만 있음":           [{"id": "d", "name": "D"}],
    "빈 목록":               [],
}
for why, ms in CASES.items():
    got = [h for h, _w, _k in wt.fixed_cols(ms)]
    assert got == WANT, "%s → 열이 다르다: %s" % (why, got)
ok += 1

# ── 2. 실제로 만든 양식에서도 자리가 같다 ────────────────────
# 두 프로젝트의 양식을 뽑아 머리글 줄을 맞대어 본다.
def heads(models):
    b = plan_form.build_year_form("시험", 2026, "2026-10", models)
    ws = load_workbook(io.BytesIO(b)).active
    # 머리글 줄은 HDR_TOP 이다 (2026-10-22 에 프로젝트 현황이
    # 위로 올라오면서 2행 → 4행으로 내려갔다)
    return [ws.cell(plan_form.HDR_TOP, i).value for i in range(1, 10)]

h1 = heads(CASES["파트넘버·유형 다 있음"])
h2 = heads(CASES["파트넘버 없음"])
assert h1 == h2, "프로젝트마다 머리글이 다르다:\n  %s\n  %s" % (h1, h2)
ok += 1

# ── 3. 그 아홉 칸이 무엇인지 못 박는다 ───────────────────────
# 순서가 바뀌면 사업부에 "A열은 파트넘버" 라고 안내한 것이 틀어진다.
# 2026-10-21 — 모델 등록 양식과 합치면서 전환주차 · 비고가 붙었다.
#   잔량은 PO−실적이라 그 둘 옆에 둔다.
# 2026-10-22 — 머리글 앞에 '✎'(적는 칸) · '='(자동 계산) 표가 붙는다.
h1 = [str(x or "").strip().strip("✎=＝※ ") for x in h1]
assert h1 == ["파트넘버", "모델명", "유형", "구분",
              "판가($)", "재료비($)", "PO수량", "실적수량", "잔량"], h1
ok += 1

# ── 4. 열두 달 주차가 다 들어 있다 ───────────────────────────
blocks = plan_form.layout(2026, "2026-10")
assert len(blocks) == 12, "달이 %d개뿐이다" % len(blocks)
assert all(b["weeks"] for b in blocks), "주차가 비어 있는 달이 있다"
assert sum(1 for b in blocks if b["open"]) == 1, "펼친 달이 하나가 아니다"
ok += 1

# ── 5. 월 합계는 적는 칸이 아니라 수식이다 ───────────────────
# 주차에 한 번 적으면 월은 따라 나온다. 두 번 적지 않게 하는 핵심이다.
b = plan_form.build_year_form("시험", 2026, "2026-10",
                              CASES["파트넘버·유형 다 있음"])
ws = load_workbook(io.BytesIO(b)).active
found = [str(c.value) for row in ws.iter_rows(min_row=plan_form.FIRST_ROW,
                                     max_row=plan_form.FIRST_ROW)
         for c in row if str(c.value or "").startswith("=")]
assert any("SUM(" in f for f in found), "월 합계가 수식이 아니다"
assert any("COUNT(" in f for f in found), "잔량 수식이 없다"
ok += 1

# ── 6. 글자 칸이 '=' 로 시작하면 안 된다 ─────────────────────
#
# 2026-10-22 — 머리글에 '적는 칸 / 자동 계산' 표를 달면서 '=' 를 썼더니
# 양식 곳곳에 #NAME? 이 떴다. 엑셀은 '=' 로 시작하는 글자를 **수식으로**
# 읽는다. '= 월 합계' 가 '월 합계 라는 이름을 계산해라' 가 되고, 그런
# 이름이 없으니 #NAME? 을 뱉는다.
#
# 이 시험이 막는 것
#   글자를 넣을 자리에 수식 기호로 시작하는 말을 쓰는 것.
_wb = load_workbook(io.BytesIO(
    plan_form.build_year_form(
        "테스트", 2026, "2026-10",
        [{"part_number": "PN%03d" % i, "name": "M%03d" % i,
          "dev_type": "자동차", "group": "양산", "price": 10,
          "material_cost": 4, "po_qty": 100, "shipped_qty": 50}
         for i in range(1, 4)],
        status_note="홀딩", hold="보류")))
_bad = []
for _sn in _wb.sheetnames:
    for _row in _wb[_sn].iter_rows():
        for _c in _row:
            _v = _c.value
            if not isinstance(_v, str) or not _v.startswith("="):
                continue
            # 진짜 수식은 함수로 시작한다
            if not _v[1:].lstrip().upper().startswith(("IF", "SUM", "COUNT")):
                _bad.append((_sn, _c.coordinate, _v[:40]))
assert not _bad, "'=' 로 시작하는 글자 칸이 있다 (엑셀이 수식으로 읽어 " \
                 "#NAME? 이 된다): %r" % _bad[:5]
ok += 1
print("6. '=' 로 시작하는 글자 칸이 없다 (#NAME? 안 뜬다) ✓")
print("전부 통과 · %d 가지" % ok)
