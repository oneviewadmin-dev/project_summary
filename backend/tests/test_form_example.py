# 빈 양식의 예시 줄은 보여 주기만 하고 올라가지 않는다
#   python3 backend/tests/test_form_example.py
#
# 처음 양식을 받는 사람은 '어느 칸에 어떤 모양으로' 를 모른다. 말로 적어
# 둔 안내보다 채워진 줄 한 개가 빠르다. 그래서 빈 양식 맨 위에 예시를
# 몇 줄 얹는다.
#
# 다만 그 줄이 진짜 모델로 올라가면 안 된다. 지우고 쓰는 게 정석이지만,
# 안 지우고 올리는 사람이 반드시 있다. 파트넘버 앞에 '(예시)' 표를 달고
# 읽는 쪽이 그 표를 보면 건너뛴다.
#
# 이 시험이 막는 것
#   예시 줄이 모델로 등록되는 것.
#   모델이 든 양식에까지 예시가 끼어드는 것.
#   예시가 '적는 줄 100개' 를 잡아먹는 것.
import ast
import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                      # noqa: E402
import pbx_plan_import as pbx                           # noqa: E402
import plan_form as pf                                  # noqa: E402


# 2026-10-22 — 머리글 앞에 '✎'(적는 칸) · '='(자동 계산) 표가 붙는다.
# 비교할 때는 떼고 본다. 읽는 쪽(_head_key · _plan_sheet_models)도 같다.
def _hk(v):
    return str(v or "").strip().strip("✎=＝*※·•-–— \t").strip()


ok = 0


def _ws(models):
    raw = pf.build_year_form("" if not models else "테스트",
                             2026, "2026-10", models)
    return load_workbook(io.BytesIO(raw))["2026-10"], raw


def _find(ws, text):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == text:
            return r
    return None


MS = [{"part_number": "PN%03d" % i, "name": "M%03d" % i, "dev_type": "자동차",
       "group": "양산", "price": 10, "material_cost": 4, "po_qty": 100,
       "shipped_qty": 50} for i in range(1, 6)]

# ── 1. 빈 양식 맨 위에 예시가 있다 ───────────────────────────
ws, _ = _ws([])
got = [str(ws.cell(r, 1).value or "")
       for r in range(pf.FIRST_ROW, pf.FIRST_ROW + len(pf.EXAMPLE))]
assert all(x.startswith(pf.EX_MARK) for x in got), "예시 줄이 없다: %r" % got
guide = str(ws.cell(pf.FIRST_ROW + len(pf.EXAMPLE), 1).value or "")
assert "여기까지 예시" in guide, "예시가 어디서 끝나는지 안 적혀 있다"
ok += 1
print("1. 빈 양식에 예시 %d 줄 + 끝 안내 ✓" % len(pf.EXAMPLE))

# ── 2. 예시는 칸을 골고루 보여 준다 ──────────────────────────
#   판가 · 수량만 채워 두면 전환주차 · 비고를 어떻게 적는지 모른다.
head = {_hk(ws.cell(pf.HDR_TOP, c).value): c for c in range(1, 14)}
cells = {k: [ws.cell(r, c).value
             for r in range(pf.FIRST_ROW, pf.FIRST_ROW + len(pf.EXAMPLE))]
         for k, c in head.items() if k}
for k in ("유형", "구분", "판가($)", "재료비($)", "PO수량", "실적수량",
          "개발→양산 전환주차", "비고"):
    assert any(v not in (None, "") for v in cells[k]), \
        "'%s' 칸을 보여 주는 예시가 없다" % k
grp = [v for v in cells["구분"] if v]
assert "양산" in grp and "개발" in grp, "양산 · 개발 둘 다 보여 줘야 한다"
ok += 1
print("2. 유형 · 구분(양산/개발) · 전환주차 · 비고까지 다 보여 준다 ✓")

# ── 3. 이번 달 주차 칸에도 숫자가 들어 있다 ──────────────────
wk = [c for c in range(1, ws.max_column + 1)
      if str(ws.cell(pf.HDR_WEEK, c).value or "").startswith("W")]
nums = [ws.cell(pf.FIRST_ROW, c).value for c in wk[:2]]
assert any(isinstance(v, (int, float)) for v in nums), \
    "주차 칸이 비어 있어 계획·실적을 어떻게 적는지 안 보인다"
ok += 1
print("3. 주차 계획 · 실적 칸도 채워져 있다 ✓")

# ── 4. 올려도 모델이 되지 않는다 ─────────────────────────────
parsed = pbx.parse_sheet(ws, "2026-10")
n = len(parsed["rows"]) if parsed else 0
assert n == 0, "예시가 %d 줄 올라간다" % n
ok += 1
print("4. 그대로 올려도 읽히는 줄 0 ✓")

# ── 5. 모델 칸 읽기도 예시를 건너뛴다 ────────────────────────
#   _plan_sheet_models 는 판가 · 재료비 · PO 를 넣는 길이다. 여기로
#   예시가 새면 '(예시) 853-...' 이라는 모델이 생긴다.
src = io.open(ROOT / "main.py", encoding="utf-8").read()
tree = ast.parse(src)
NS = {}
for node in tree.body:
    if isinstance(node, ast.Assign) and getattr(
            node.targets[0], "id", "") in ("_PLAN_MODEL_COLS", "_PLAN_STOP"):
        exec(compile(ast.Module([node], []), "<m>", "exec"), NS)
    if isinstance(node, ast.FunctionDef) and node.name in (
            "_plan_sheet_models", "_as_int", "_as_money"):
        exec(compile(ast.Module([node], []), "<m>", "exec"), NS)
rows = NS["_plan_sheet_models"](ws)
assert not rows, "모델 칸 읽기로 예시 %d 줄이 샜다: %r" % (
    len(rows), [r["name"] for r in rows[:3]])
ok += 1
print("5. 모델 칸 읽기도 예시를 건너뛴다 ✓")

# ── 6. 모델이 든 양식에는 예시를 안 넣는다 ───────────────────
ws2, _ = _ws(MS)
first = str(ws2.cell(pf.FIRST_ROW, 1).value or "")
assert not first.startswith(pf.EX_MARK), "모델 양식에 예시가 끼었다"
ok += 1
print("6. 모델이 든 양식에는 예시가 없다 ✓")

# ── 7. 예시가 적는 줄 100개를 잡아먹지 않는다 ────────────────
r_sum = _find(ws, "합계")
rows_n = r_sum - pf.FIRST_ROW - len(pf.EXAMPLE) - 1      # 예시 + 끝 안내
assert rows_n == pf.TOTAL_ROWS, "적는 줄이 %d 줄이다" % rows_n
ok += 1
print("7. 예시 아래로 적는 줄이 %d 줄 그대로 ✓" % rows_n)

print("전부 통과 · %d 가지" % ok)
