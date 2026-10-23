# 모델 열을 고를 때 '더 많이 차 있는 쪽' 을 고른다
#   python3 backend/tests/test_model_col_pick.py
#
# 2026-10-22 — 챔버 양식을 올리면 24종 중 1종만 들어갔다. 아무 말도 안
# 나왔다. 까닭은 이랬다.
#   · 챔버는 모델 24종 중 파트넘버가 적힌 모델이 하나뿐이다
#     (나머지는 품번이 모델명 칸에 들어가 있다).
#   · 읽는 쪽은 "파트넘버 열에 값이 하나라도 있으면 그 열" 로 골랐다.
#   · 그 하나 때문에 파트넘버 열이 뽑혔고, 라벨이 빈 23줄은 버려졌다.
#
# 이 시험이 막는 것
#   한 칸 때문에 거의 빈 열을 고르는 것.
#   세는 범위가 '합계' 줄을 넘어 아래 표까지 내려가 열을 뒤바꾸는 것.
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook                           # noqa: E402
import pbx_plan_import as pbx                           # noqa: E402

ok = 0


def _sheet(pairs, tail=None):
    """pairs = [(파트넘버, 모델명), ...] 로 작은 양식 하나를 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.cell(2, 1, "파트넘버")
    ws.cell(2, 2, "모델명")
    ws.cell(3, 5, "W41")
    ws.cell(4, 5, "계획")
    ws.cell(4, 6, "실적")
    r = 5
    for pn, nm in pairs:
        if pn:
            ws.cell(r, 1, pn)
        if nm:
            ws.cell(r, 2, nm)
        ws.cell(r, 5, 1)
        r += 1
    ws.cell(r, 1, "합계")                   # 읽기가 여기서 끝난다
    for extra in (tail or []):             # 그 아래 붙은 표
        r += 2
        ws.cell(r, 1, extra[0])
        ws.cell(r, 2, extra[1])
    return ws


# ── 1. 파트넘버가 하나뿐이면 모델명 열을 쓴다 (챔버) ─────────
ws = _sheet([(None, "715-A39794-%03d" % i) for i in range(1, 24)]
            + [("853-268706-003", "853-268706-003")])
c = pbx._find_model_col(ws, 4, 5)
assert c == 2, "모델명 열(2)을 골라야 한다 — 고른 열 %s" % c
got = pbx._rows_of(ws, 4, c, "2026-10", {41: (5, 6)})
assert len(got["rows"]) == 24, \
    "24줄이 다 읽혀야 한다 — 읽은 줄 %d" % len(got["rows"])
ok += 1
print("1. 파트넘버가 한 칸뿐이어도 24줄 전부 읽힌다 ✓")

# ── 2. 파트넘버가 다 차 있으면 파트넘버 열 (파워박스) ────────
ws = _sheet([("853-800575-%03d" % i, "PBX-%03d" % i) for i in range(1, 31)])
assert pbx._find_model_col(ws, 4, 5) == 1, "파트넘버 열(1)을 골라야 한다"
ok += 1
print("2. 양쪽 다 차 있으면 파트넘버 열 ✓")

# ── 3. 둘이 같은 수면 파트넘버가 이긴다 ──────────────────────
ws = _sheet([("A%d" % i, "A%d" % i) for i in range(1, 11)])
assert pbx._find_model_col(ws, 4, 5) == 1, "같으면 파트넘버 열이어야 한다"
ok += 1
print("3. 같은 수면 파트넘버 ✓")

# ── 4. '합계' 아래 붙은 표는 세지 않는다 ─────────────────────
#   적는 칸 밑의 'PPT 붙이는 표' 는 1열이 구분, 2열이 파트넘버다. 거기까지
#   세면 모델명 열이 두 배로 보여서 열이 뒤바뀐다.
ws = _sheet([("853-800575-%03d" % i, "PBX-%03d" % i) for i in range(1, 31)],
            tail=[("양산", "853-800575-%03d" % i) for i in range(1, 31)])
c = pbx._find_model_col(ws, 4, 5)
assert c == 1, "아래 표까지 세어 모델명 열(%s)을 골랐다" % c
ok += 1
print("4. '합계' 아래 표는 세지 않는다 ✓")

# ── 5. 파트넘버 칸이 통째로 빈 프로젝트 (CUP · 하바플레이트) ──
ws = _sheet([(None, "02-434103-%02d" % i) for i in range(1, 30)])
assert pbx._find_model_col(ws, 4, 5) == 2, "모델명 열이어야 한다"
ok += 1
print("5. 파트넘버가 통째로 비면 모델명 열 ✓")

print("전부 통과 · %d 가지" % ok)
