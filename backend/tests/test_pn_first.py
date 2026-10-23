# 파트넘버가 먼저고, 모델명은 다를 때만 적는다
#   python3 backend/tests/test_pn_first.py
#
# 2026-10-22 — 챔버 양식에 853-268706-003 이 파트넘버 칸과 모델명 칸에
# 똑같이 두 번 찍혀 나왔다. 등록된 자료가 세 모양으로 섞여 있어서다.
#   · 두 칸이 제대로 나뉨            853-800575-413 / AETHER GDX
#   · 두 칸에 같은 값               853-268706-003 / 853-268706-003   (파워박스 44종)
#   · 파트넘버가 비고 품번이 모델명에  (빈 칸)      / 715-A39794-613   (챔버 23종)
#
# 규칙 하나로 정리한다. 품번은 늘 파트넘버 칸. 모델명은 품번과 **다를 때만**.
#
# 이 시험이 막는 것
#   같은 글자가 두 칸에 다시 찍히는 것.
#   품번이 모델명 칸에만 있는 프로젝트에서 파트넘버 칸이 비는 것.
#   읽는 쪽(_plan_sheet_models)이 모델명 칸만 보고 그런 줄을 버리는 것.
import ast
import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                       # noqa: E402
import week_calendar as wcal                            # noqa: E402
import weekly_template as wt                             # noqa: E402

ok = 0

CASES = [
    # (모델, 기대 파트넘버, 기대 모델명)
    ({"id": "a", "part_number": "853-800575-413", "name": "AETHER GDX"},
     "853-800575-413", "AETHER GDX"),
    ({"id": "b", "part_number": "853-268706-003", "name": "853-268706-003"},
     "853-268706-003", ""),
    ({"id": "c", "name": "715-A39794-613"},
     "715-A39794-613", ""),
    ({"id": "715-156359-006"},
     "715-156359-006", ""),
    # 대소문자만 다른 것도 같은 값으로 본다
    ({"id": "e", "part_number": "ab-100", "name": "AB-100"}, "ab-100", ""),
]

# ── 1. 세 모양 모두 규칙대로 갈린다 ──────────────────────────
for m, want_pn, want_nm in CASES:
    pn, nm = wt._pn_and_name(m)
    assert (pn, nm) == (want_pn, want_nm), \
        "%r → (%r, %r) · (%r, %r) 여야 한다" % (m, pn, nm, want_pn, want_nm)
ok += 1
print("1. 파트넘버 우선 · 같으면 모델명 비움 — %d 가지 ✓" % len(CASES))

# ── 2. 실제 양식에도 그대로 나온다 ───────────────────────────
ms = [dict(m, dev_type="자동차", group="양산") for m, _, _ in CASES]
raw = wt.build_weekly_template("T", sheets=[{"month": "2026-10",
                                      "weeks": wcal.get_month_weeks("2026-10")}], models=ms)
ws = load_workbook(io.BytesIO(raw))["2026-10"]
head = {str(ws.cell(2, c).value or "").strip(): c for c in range(1, 12)}
got = []
for r in range(wt.FIRST_ROW, wt.FIRST_ROW + len(ms)):
    got.append((str(ws.cell(r, head["파트넘버"]).value or "").strip(),
                str(ws.cell(r, head["모델명"]).value or "").strip()))
assert all(pn for pn, _ in got), "파트넘버가 빈 줄이 있다: %r" % got
assert not [g for g in got if g[0] and g[0].lower() == g[1].lower()], \
    "같은 글자가 두 칸에 찍혔다: %r" % got
ok += 1
print("2. 양식에 파트넘버가 다 차고 같은 글자가 두 번 안 찍힌다 ✓")

# ── 3. 읽는 쪽이 모델명 빈 줄을 안 버린다 ────────────────────
#   _plan_sheet_models 는 main.py 안에 있다. 서버를 띄우지 않고 그 함수만
#   꺼내 쓴다 — main.py 를 import 하면 서버가 통째로 올라온다.
src = io.open(ROOT / "main.py", encoding="utf-8").read()
tree = ast.parse(src)
NS = {}
WANT = {"_plan_sheet_models", "_as_int", "_as_money"}
for node in tree.body:
    if isinstance(node, ast.Assign) and getattr(
            node.targets[0], "id", "") in ("_PLAN_MODEL_COLS", "_PLAN_STOP"):
        exec(compile(ast.Module([node], []), "<main>", "exec"), NS)
    if isinstance(node, ast.FunctionDef) and node.name in WANT:
        exec(compile(ast.Module([node], []), "<main>", "exec"), NS)
rows = NS["_plan_sheet_models"](ws)
assert len(rows) == len(ms), \
    "모델 %d종 중 %d줄만 읽혔다 (모델명이 빈 줄을 버렸다)" % (len(ms), len(rows))
names = [r["name"] for r in rows]
assert names == [pn for pn, _ in got], \
    "읽은 이름이 파트넘버와 다르다: %r" % names
ok += 1
print("3. 읽는 쪽이 %d줄 전부 읽고, 이름을 파트넘버로 잡는다 ✓" % len(rows))

print("전부 통과 · %d 가지" % ok)
