# 구분(양산/개발)과 전환 주차는 엑셀로만 정한다
#   python3 backend/tests/test_group_from_excel.py
#
# 전에는 관리자 화면에서 구분을 고르고, 전환 주차도 손으로 넣었다. 엑셀과
# 어긋나면 어느 쪽이 맞는지 알 길이 없었고, 매출은 양산(판가 x 수량)과
# 개발($3,400 고정)이 달라 숫자까지 흔들렸다. 그래서 기준을 엑셀 하나로 모았다.
#
#   모델 등록 엑셀  구분 칸        → 지금 구분
#                   전환주차 칸    → 개발 → 양산으로 넘어간 주차
#   관리자 화면                    → 보기만 한다
import ast
import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))
ok = 0

# ── 1. 양식에 전환주차 칸이 있다 ─────────────────────────────
import openpyxl                                          # noqa: E402
import model_template                                    # noqa: E402

wb = openpyxl.load_workbook(io.BytesIO(
    model_template.build_model_template("시험", ["413"])))
ws = wb["모델 목록"]
heads = [ws.cell(2, i).value for i in range(1, ws.max_column + 1)]
assert "개발→양산 전환주차" in heads, "양식에 전환주차 칸이 없다: %s" % heads
ok += 1

# 작성 안내에도 설명이 있다 — 칸만 있고 설명이 없으면 아무도 안 적는다
guide = "\n".join(str(c.value or "") for row in wb["작성 안내"].iter_rows() for c in row)
assert "개발→양산 전환주차" in guide and "2026-W36" in guide, "작성 안내에 전환주차 설명이 없다"
ok += 1

# ── 2. 임포터가 전환주차 칸을 알아본다 ───────────────────────
SRC = (ROOT / "main.py").read_text(encoding="utf-8")
assert 'col_map["switch_week"]' in SRC, "임포터가 전환주차 열을 안 읽는다"
# 임포터는 머리글에서 띄어쓰기를 떼고 본다 — 그 꼴로 들어 있어야 한다
for word in ("개발→양산전환주차", "양산전환"):
    assert word in SRC, "머리글 '%s' 를 못 알아본다" % word
ok += 1

# ── 3. 전환주차 한 칸이 구간 두 개로 펴진다 ──────────────────
assert "_phases_from_switch" in SRC, "전환주차를 phases 로 펴는 곳이 없다"
assert SRC.count("_phases_from_switch") >= 3, \
    "새 모델과 기존 모델 양쪽에 적용되지 않는다"
ok += 1

# '-' 로 지울 수 있다 (처음부터 한 구분이라는 뜻)
assert 'existing.pop("phases", None)' in SRC, "전환 이력을 지우는 길이 없다"
ok += 1

# ── 4. 화면에서는 못 고친다 ──────────────────────────────────
HTML = (ROOT / "admin_v2.html").read_text(encoding="utf-8")
assert 'data-field="group" disabled' in HTML, "모델 목록의 구분이 아직 고를 수 있다"
assert 'class="ph-group" disabled' in HTML, "이력 편집기의 구분이 아직 고를 수 있다"
assert "ph-from" in HTML and "readonly" in HTML, "전환 주차 칸이 아직 고칠 수 있다"
assert "ph-add" not in HTML, "구간 추가 버튼이 남아 있다 — 전환을 손으로 만들 수 있다"
assert "ph-del" not in HTML, "구간 삭제 버튼이 남아 있다"
ok += 1

# 어떻게 바꾸는지가 화면에 적혀 있다
assert ("전환주차" in HTML), "엑셀로 바꾸라는 안내가 화면에 없다"
ok += 1

# ── 5. 판가·재료비는 여전히 고칠 수 있다 ─────────────────────
# 판가는 계약이 바뀌면 진짜로 바뀐다. 여기까지 막으면 쓸 수 없게 된다.
assert 'class="ph-price" type="number"' in HTML, "판가까지 막혔다"
assert 'class="ph-mc" type="number"' in HTML, "재료비까지 막혔다"
ok += 1

print("전부 통과 ·", ok, "가지")
