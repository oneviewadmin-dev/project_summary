# 프로젝트 보류/드롭과 모델 보류/드롭은 다른 층이다
#   python3 backend/tests/test_proj_vs_model_hold.py
#
# 보류 · 드롭예정이 두 층으로 있다.
#   · 프로젝트째 멈춘 것   CUP "컵 이슈 해결될 때까지 LAIR 작성 홀딩"
#     → 그 프로젝트 모델 전부의 시계가 멈춘다.
#   · 모델 하나만 멈춘 것   파워박스 VCTR-XPRSMS "드롭 예정"
#     → 그 한 줄만 멈춘다.
#
# 섞이면 어느 쪽이든 숫자가 거짓말을 한다. 모델 하나 드롭을 프로젝트
# 드롭으로 읽으면 멀쩡한 모델 수십 종이 지연 집계에서 통째로 빠지고,
# 반대면 멈춰 세운 프로젝트가 전부 '지연' 으로 세어진다.
#
# 이 시험이 막는 것
#   양식에서 두 칸이 한 칸으로 합쳐지는 것.
#   모델 비고의 '드롭' 이 프로젝트 현황으로 새는 것(그 반대도).
import ast
import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]      # backend/
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                      # noqa: E402
import pbx_plan_import as pbx                           # noqa: E402
import plan_form as pf                                  # noqa: E402


# 2026-10-22 — 머리글 앞에 '✎'(적는 칸) · '='(자동 계산) 표가 붙는다.
# 비교할 때는 떼고 본다. 읽는 쪽(_head_key · _plan_sheet_models)도 같다.
def _hk(v):
    return str(v or "").strip().strip("✎=＝*※·•-–— \t").strip()


ok = 0

MS = [{"part_number": "PN%03d" % i, "name": "M%03d" % i, "dev_type": "자동차",
       "group": "양산", "price": 10, "material_cost": 4, "po_qty": 100,
       "shipped_qty": 50, "note": ""} for i in range(1, 4)]


def _sheet(models=None, note="", hold=""):
    raw = pf.build_year_form("테스트", 2026, "2026-10", models or MS,
                             status_note=note, hold=hold)
    return load_workbook(io.BytesIO(raw))["2026-10"], raw


def _find(ws, text):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == text:
            return r
    return None


# ── 1. 양식에 두 칸이 따로 있다 ──────────────────────────────
ws, _ = _sheet()
r_proj = _find(ws, pf.PROJ_LABEL)
assert r_proj, "'프로젝트 현황' 칸이 없다"
head = {_hk(ws.cell(pf.HDR_TOP, c).value): c for c in range(1, 14)}
assert "비고" in head, "모델 줄의 비고 칸이 없다"
# 2026-10-22 — 현황은 **적는 칸 위**다 (제목 바로 밑).
assert r_proj == pf.ROW_PROJ, "프로젝트 현황이 %s 줄에 있다" % r_proj
assert r_proj < pf.FIRST_ROW, "프로젝트 현황이 적는 칸 안에 있다"
ok += 1
print("1. 모델 비고(%s열) 와 프로젝트 현황(%d줄) 이 따로 있다 ✓"
      % (head["비고"], r_proj))

# ── 2. 프로젝트 현황을 적으면 프로젝트만 멈춘다 ──────────────
for hold, note in (("보류", "컵 이슈 해결될 때까지 홀딩"),
                   ("드롭예정", "고객사 드롭 통보"),
                   ("", "정상 진행 중")):
    ws, _ = _sheet(note=note, hold=hold)
    got = pbx.read_project_status(ws)
    assert (got["hold"], got["note"]) == (hold, note), \
        "%r → %r" % ((hold, note), got)
    assert got["blank"] is False, "고른 값이 '빈 칸' 으로 읽혔다: %r" % got
ok += 1
print("2. 프로젝트 현황을 그대로 읽는다 — 보류 · 드롭예정 · 진행 ✓")

# ── 2-2. 상태 칸이 비면 '그대로 둬라' ────────────────────────
#   빈 칸과 '진행' 은 다르다. 빈 칸은 '안 건드렸다', '진행' 은 '풀어 달라'.
#   섞으면 상태 칸을 안 본 사람이 올릴 때마다 멈춰 세운 프로젝트가 풀린다.
ws, _ = _sheet(note="컵 이슈 홀딩", hold="보류")
r = _find(ws, pf.PROJ_LABEL)
ws.cell(r + 1, 2).value = None               # 사람이 상태 칸을 비웠다
got = pbx.read_project_status(ws)
assert got["blank"] is True, "빈 칸을 '진행' 으로 읽었다: %r" % got
assert got["note"] == "컵 이슈 홀딩", "글은 그대로 읽혀야 한다"
ok += 1
print("2-2. 상태 칸이 비면 '그대로 둬라' 로 읽는다 ✓")

# ── 3. 모델 비고의 '드롭' 은 프로젝트로 새지 않는다 ──────────
ms = [dict(MS[0], note="드롭 예정"), MS[1], MS[2]]
ws, _ = _sheet(models=ms)
got = pbx.read_project_status(ws)
assert got["hold"] == "" and got["note"] == "", \
    "모델 비고가 프로젝트 현황으로 샜다: %r" % got
# 양식은 지금 상태를 미리 채워서 나간다(멈춘 데가 없으면 '진행').
# 그래서 손 안 댄 양식은 빈 칸이 아니라 '진행' 을 고른 것으로 읽힌다.
assert got["blank"] is False, "미리 채운 '진행' 이 빈 칸으로 읽혔다: %r" % got
ok += 1
print("3. 모델 비고의 '드롭 예정' 은 프로젝트 현황에 안 들어간다 ✓")

# ── 4. 프로젝트 현황 글은 모델 비고로 새지 않는다 ────────────
ws, _ = _sheet(note="컵 이슈 해결될 때까지 홀딩", hold="보류")
nb = head["비고"]
notes = [str(ws.cell(r, nb).value or "").strip()
         for r in range(pf.FIRST_ROW, pf.FIRST_ROW + len(MS))]
assert not any(notes), "프로젝트 현황이 모델 비고로 샜다: %r" % notes
ok += 1
print("4. 프로젝트 현황 글이 모델 비고로 안 들어간다 ✓")

# ── 5. 서버가 두 층을 따로 읽는다 ────────────────────────────
#   _model_hold 는 모델의 note/issues 만, _project_hold 는 프로젝트의
#   hold/status_note 만 본다. 한쪽이 다른 쪽을 보기 시작하면 터진다.
src = io.open(ROOT / "main.py", encoding="utf-8").read()
tree = ast.parse(src)


def body(name):
    for n in tree.body:
        if isinstance(n, ast.FunctionDef) and n.name == name:
            return ast.get_source_segment(src, n)
    return ""


mh, ph = body("_model_hold"), body("_project_hold")
assert mh and ph, "_model_hold · _project_hold 가 없다"
assert "status_note" not in mh, "_model_hold 가 프로젝트 현황을 본다"
assert '"note"' in mh and '"issues"' in mh, "_model_hold 가 비고를 안 본다"
assert "status_note" in ph, "_project_hold 가 프로젝트 현황을 안 본다"
ok += 1
print("5. 서버의 _model_hold · _project_hold 가 서로 안 넘본다 ✓")

# ── 6. 올릴 때 프로젝트 현황을 반영한다 ──────────────────────
i = src.find("프로젝트 현황 — 양식 아래")
assert i > 0, "올리는 자리에서 프로젝트 현황을 안 읽는다"
seg = src[i:i + 1400]
assert "read_project_status" in seg, "read_project_status 를 안 부른다"
assert 'proj["status_note"]' in seg, "현황 글을 안 넣는다"
assert '"excel"' in seg, "엑셀에서 왔다는 표시를 안 남긴다"
assert '_pstat["blank"]' in seg, "빈 칸을 '그대로 둬라' 로 안 본다"
ok += 1
print("6. 올리면 프로젝트 현황이 반영된다 ✓")

# ── 7. 엑셀로 고른 값을 현황 문구가 덮지 않는다 ──────────────
#   _stamp_hold_on_boot 은 서버가 뜰 때마다 돈다. 엑셀에서 '진행' 으로
#   바꿔 올려도 현황 글에 '홀딩' 이라는 말이 남아 있으면 다음 배포에서
#   조용히 보류로 되돌아갔다.
st = body("_stamp_project_hold")
assert st, "_stamp_project_hold 가 없다"
assert '("manual", "excel")' in st, \
    "엑셀에서 고른 값을 현황 문구가 덮는다 — 배포할 때마다 되돌아간다"
ok += 1
print("7. 엑셀로 고른 값이 배포 때 되돌아가지 않는다 ✓")

print("전부 통과 · %d 가지" % ok)
